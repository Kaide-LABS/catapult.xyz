import shutil
import openpyxl
from models.schemas import VSQState

def export_node(state: VSQState):
    original_path = state["file_path"]
    export_path = original_path.replace(".xlsx", "_completed.xlsx").replace(".csv", "_completed.xlsx")
    
    log = []
    
    if state["file_type"] == "xlsx":
        try:
            shutil.copy(original_path, export_path)
            wb = openpyxl.load_workbook(export_path)
            sheet = wb.active
            
            ans_col = sheet.max_column + 1
            sheet.cell(row=1, column=ans_col, value="Answer")
            sheet.cell(row=1, column=ans_col+1, value="Confidence")
            sheet.cell(row=1, column=ans_col+2, value="Citations")
            
            ans_map = {}
            for q_id, approved in state.get("approved_answers", {}).items():
                q_text = ""
                for q in state.get("questions", []):
                    if q.id == q_id:
                        q_text = q.original_text
                        break
                if q_text:
                    ans_map[q_text] = approved
                    
            for row in range(2, sheet.max_row + 1):
                for col in range(1, sheet.max_column):
                    cell_val = str(sheet.cell(row=row, column=col).value).strip()
                    if cell_val in ans_map:
                        approved = ans_map[cell_val]
                        sheet.cell(row=row, column=ans_col, value=approved.final_answer)
                        sheet.cell(row=row, column=ans_col+1, value=approved.confidence_score)
                        cites = "\\n".join([f"{c.source_document} ({c.section})" for c in approved.citations])
                        sheet.cell(row=row, column=ans_col+2, value=cites)
                        break
            
            audit_ws = wb.create_sheet("Audit Trail")
            audit_ws.append(["Question ID", "Source Document", "Section", "Chunk ID", "Verbatim Quote"])
            for q_id, approved in state.get("approved_answers", {}).items():
                for c in approved.citations:
                    audit_ws.append([q_id, c.source_document, c.section, c.chunk_id, c.quote])
                    
            wb.save(export_path)
            log.append(f"[Export] Saved completed Excel to {export_path}")
        except Exception as e:
            log.append(f"[Export] Excel export failed: {str(e)}")
            export_path = ""
    else:
        import pandas as pd
        data = []
        for q_id, approved in state.get("approved_answers", {}).items():
            q_text = next((q.original_text for q in state.get("questions", []) if q.id == q_id), "")
            cites = "\\n".join([f"{c.source_document} ({c.section})" for c in approved.citations])
            data.append({
                "ID": q_id,
                "Question": q_text,
                "Answer": approved.final_answer,
                "Confidence": approved.confidence_score,
                "Citations": cites
            })
        df = pd.DataFrame(data)
        df.to_excel(export_path, index=False)
        log.append(f"[Export] Saved fallback Excel to {export_path}")
        
    return {"export_path": export_path, "processing_log": log}
